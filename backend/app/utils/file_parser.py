"""
通用文件解析器 — PDF / DOCX / TXT / MD / XLSX
纯工具层: 输入文件字节流, 输出纯文本。不涉及数据库 / 外部服务调用。

解析失败抛 AppException 子类, 由向量化流水线捕获并置任务 failed。
"""

import io
import re
from collections.abc import Callable

from loguru import logger

from app.utils.exceptions import (
    AppException,
    BadRequestException,
    UnsupportedFileTypeException,
)

# 解析器支持的文件类型 (与 config.UPLOAD_ALLOWED_EXTENSIONS 对齐)
SUPPORTED_FILE_TYPES = ("pdf", "docx", "txt", "md", "xlsx")

# 编码回退顺序 (txt/md)
_TEXT_ENCODINGS = ("utf-8", "gbk", "latin-1")


# ==========================================
# 各类型解析实现
# ==========================================


def _parse_txt(file_bytes: bytes) -> str:
    """TXT: 多编码尝试解码"""
    for encoding in _TEXT_ENCODINGS:
        try:
            return file_bytes.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return file_bytes.decode("utf-8", errors="ignore")


def _parse_md(file_bytes: bytes) -> str:
    """Markdown: 解码后清理常用语法符号, 保留正文语义"""
    text = _parse_txt(file_bytes)
    # 图片语法 ![...](...) → 仅保留描述文字
    text = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", text)
    # 链接语法 [...](...) → 仅保留文字
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    # 清理标题 / 加粗 / 行内代码等标记符号
    text = re.sub(r"^#{1,6}\s+", "", text, flags=re.MULTILINE)
    text = re.sub(r"(\*\*|__|`)(.*?)\1", r"\2", text)
    return text


def _parse_docx(file_bytes: bytes) -> str:
    """DOCX: 段落 + 表格文本 (python-docx)"""
    from docx import Document

    doc = Document(io.BytesIO(file_bytes))
    parts: list[str] = []

    for para in doc.paragraphs:
        line = para.text.strip()
        if line:
            parts.append(line)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def _parse_xlsx(file_bytes: bytes) -> str:
    """XLSX: 遍历所有工作表, 单元格制表符拼接 (openpyxl)"""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    parts: list[str] = []

    for ws in wb.worksheets:
        parts.append(f"# 工作表: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if cell is None else str(cell).strip() for cell in row]
            if any(cells):
                parts.append("\t".join(cells))
    wb.close()

    return "\n".join(parts)


def _parse_pdf(file_bytes: bytes) -> str:
    """PDF: pdfplumber 优先, 失败回退 PyPDF2"""
    errors: list[str] = []

    try:
        import pdfplumber

        parts: list[str] = []
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    parts.append(text)
        return "\n".join(parts)
    except Exception as e:  # pdfplumber 解析失败
        errors.append(f"pdfplumber: {e}")

    try:
        from PyPDF2 import PdfReader

        reader = PdfReader(io.BytesIO(file_bytes))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as e:  # PyPDF2 亦失败
        errors.append(f"PyPDF2: {e}")

    raise BadRequestException(f"PDF 解析失败: {'; '.join(errors)}")


# 解析器注册表
_PARSERS: dict[str, Callable[[bytes], str]] = {
    "txt": _parse_txt,
    "md": _parse_md,
    "docx": _parse_docx,
    "xlsx": _parse_xlsx,
    "pdf": _parse_pdf,
}


# ==========================================
# 统一入口
# ==========================================


def parse_file(file_bytes: bytes, file_type: str) -> str:
    """
    按文件类型解析字节流为纯文本。

    Args:
        file_bytes: 文件字节内容
        file_type: 文件后缀 (小写, 如 "pdf")

    Returns:
        解析后的纯文本 (已 strip)

    Raises:
        UnsupportedFileTypeException: 不支持的文件类型
        BadRequestException: 文件内容解析失败
    """
    file_type = (file_type or "").lower()

    parser = _PARSERS.get(file_type)
    if parser is None:
        raise UnsupportedFileTypeException(
            f"不支持的文件类型 '.{file_type}', 支持: {', '.join(SUPPORTED_FILE_TYPES)}"
        )

    if not file_bytes:
        return ""

    try:
        text = parser(file_bytes)
    except AppException:  # 业务异常 (如 PDF 解析失败) 直接透传
        raise
    except Exception as e:
        logger.warning(f"[解析] {file_type} 文件解析失败: {e}")
        raise BadRequestException(f"文件解析失败: {e}") from e

    if not text or not text.strip():
        raise BadRequestException("文档解析结果为空，请检查文件内容")

    return text.strip()
