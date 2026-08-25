"""
通用文件解析器测试
覆盖: TXT 多编码 / Markdown 清理 / DOCX 段落与表格 / XLSX 多工作表 /
      不支持类型 / 空内容 / 解析失败
"""

import io

import pytest

from app.utils.exceptions import BadRequestException, UnsupportedFileTypeException
from app.utils.file_parser import parse_file


class TestTxt:
    def test_utf8(self):
        text = parse_file("企业知识库智能助手".encode(), "txt")
        assert text == "企业知识库智能助手"

    def test_gbk_fallback(self):
        text = parse_file("中文文档内容".encode("gbk"), "txt")
        assert text == "中文文档内容"


class TestMarkdown:
    def test_markdown_syntax_cleanup(self):
        md = (
            "# 标题\n\n"
            "这是**加粗**文本，包含[链接](http://example.com)。\n\n"
            "![图片描述](http://example.com/img.png)\n\n"
            "正文内容"
        )
        text = parse_file(md.encode("utf-8"), "md")
        assert "#" not in text
        assert "**" not in text
        assert "http://example.com" not in text
        assert "标题" in text
        assert "链接" in text
        assert "图片描述" in text
        assert "正文内容" in text


class TestDocx:
    def test_paragraphs_and_tables(self):
        from docx import Document

        doc = Document()
        doc.add_paragraph("第一段内容")
        doc.add_paragraph("第二段内容")
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "姓名"
        table.rows[0].cells[1].text = "部门"
        table.rows[1].cells[0].text = "张三"
        table.rows[1].cells[1].text = "技术部"

        buf = io.BytesIO()
        doc.save(buf)

        text = parse_file(buf.getvalue(), "docx")
        assert "第一段内容" in text
        assert "第二段内容" in text
        assert "姓名" in text and "部门" in text
        assert "张三" in text and "技术部" in text


class TestXlsx:
    def test_multiple_sheets(self):
        from openpyxl import Workbook

        wb = Workbook()
        ws1 = wb.active
        ws1.title = "员工"
        ws1.append(["姓名", "部门"])
        ws1.append(["张三", "技术部"])
        ws2 = wb.create_sheet("制度")
        ws2.append(["标题", "内容"])
        ws2.append(["考勤制度", "每日 9 点上班"])

        buf = io.BytesIO()
        wb.save(buf)

        text = parse_file(buf.getvalue(), "xlsx")
        assert "工作表: 员工" in text
        assert "张三" in text
        assert "工作表: 制度" in text
        assert "考勤制度" in text


class TestErrorCases:
    def test_unsupported_type(self):
        with pytest.raises(UnsupportedFileTypeException):
            parse_file(b"dummy", "exe")

    def test_unknown_type(self):
        with pytest.raises(UnsupportedFileTypeException):
            parse_file(b"dummy", "doc")

    def test_empty_bytes_returns_empty(self):
        assert parse_file(b"", "txt") == ""

    def test_whitespace_only_content(self):
        with pytest.raises(BadRequestException):
            parse_file(b"   \n\n  ", "txt")

    def test_invalid_docx_content(self):
        """非 DOCX 字节流 → 解析失败"""
        with pytest.raises(BadRequestException):
            parse_file(b"this is not a docx file", "docx")

    def test_invalid_xlsx_content(self):
        with pytest.raises(BadRequestException):
            parse_file(b"this is not an xlsx file", "xlsx")
